import os
import re
from collections import Counter
from io import BytesIO

import streamlit as st
import pdfplumber
from docx import Document

skills_by_category = {
    "Core Programming": [
        "python",
        "java",
        "sql",
        "javascript",
        "c++",
        "c#",
    ],
    "Web Development": [
        "html",
        "css",
        "react",
        "node.js",
        "flask",
        "django",
        "fastapi",
    ],
    "Data & AI": [
        "machine learning",
        "deep learning",
        "data analysis",
        "pandas",
        "numpy",
        "scikit-learn",
        "pytorch",
        "tensorflow",
        "nlp",
    ],
    "Tools & Workflow": [
        "git",
        "github",
        "docker",
        "linux",
        "aws",
        "azure",
        "tableau",
        "agile",
        "jira",
    ],
}

all_skills = [skill for skills in skills_by_category.values() for skill in skills]
common_stopwords = {
    "the",
    "and",
    "for",
    "with",
    "from",
    "into",
    "this",
    "that",
    "have",
    "your",
    "resume",
    "skills",
    "experience",
    "work",
    "project",
    "projects",
    "years",
    "year",
    "role",
    "based",
    "using",
    "used",
    "strong",
    "detail",
    "team",
    "also",
}


def normalize_text(text):
    text = text.lower()
    text = text.replace("\u2013", " ").replace("\u2014", " ")
    text = re.sub(r"[^a-z0-9+.#\s]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def contains_skill(text, skill):
    pattern = r"\b" + re.escape(skill) + r"\b"
    return bool(re.search(pattern, text))


def _read_file_bytes(file):
    if isinstance(file, (str, os.PathLike)):
        with open(file, "rb") as f:
            return f.read()

    if hasattr(file, "read"):
        current_position = file.tell() if hasattr(file, "tell") else None
        file.seek(0)
        data = file.read()
        if current_position is not None and hasattr(file, "seek"):
            file.seek(current_position)
        return data

    return file


def read_pdf(file):
    pdf_bytes = _read_file_bytes(file)
    text = ""
    with pdfplumber.open(BytesIO(pdf_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            if page_text:
                text += page_text + "\n"
    return text


def read_docx(file):
    docx_bytes = _read_file_bytes(file)
    doc = Document(BytesIO(docx_bytes))
    paragraphs = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            paragraphs.append(" ".join(cell.text for cell in row.cells))
    return "\n".join(paragraphs)


def extract_top_keywords(text, limit=10):
    cleaned = normalize_text(text)
    words = re.findall(r"\b[a-z0-9+#.]{2,}\b", cleaned)
    counts = Counter(word for word in words if word not in common_stopwords)
    return counts.most_common(limit)


def summarize_resume(text):
    normalized = normalize_text(text)
    word_count = len(normalized.split())
    email_present = bool(re.search(r"\b[a-z0-9._%+-]+@[a-z0-9.-]+\.[a-z]{2,}\b", text, re.I))
    phone_present = bool(re.search(r"\b(?:\+?\d{1,3}[-.\s]?)?(?:\(?\d{3}\)?[-.\s]?)\d{3}[-.\s]?\d{4}\b", text))
    return {
        "word_count": word_count,
        "email_present": email_present,
        "phone_present": phone_present,
    }


st.set_page_config(page_title="AI Resume Analyzer", page_icon="📄", layout="wide")
st.title("AI Resume Analyzer")

uploaded_file = st.file_uploader("Upload Resume", type=["pdf", "docx"])

if uploaded_file is not None:
    file_name = (uploaded_file.name or "").lower()
    if file_name.endswith(".pdf"):
        resume_text = read_pdf(uploaded_file)
    else:
        resume_text = read_docx(uploaded_file)

    normalized_resume = normalize_text(resume_text)
    found_by_category = {}
    missing_by_category = {}
    found_overall = []

    for category, skills in skills_by_category.items():
        category_found = []
        for skill in skills:
            if contains_skill(normalized_resume, skill):
                category_found.append(skill)
        found_by_category[category] = category_found
        missing_by_category[category] = [skill for skill in skills if skill not in category_found]
        found_overall.extend(category_found)

    found_overall = sorted(set(found_overall))
    missing_overall = [skill for skill in all_skills if skill not in found_overall]
    score = int((len(found_overall) / len(all_skills)) * 100) if all_skills else 0

    summary = summarize_resume(resume_text)
    top_keywords = extract_top_keywords(resume_text)

    st.subheader("ATS Score")
    col1, col2, col3 = st.columns([2, 1, 1])
    with col1:
        st.progress(score / 100)
        st.metric("Overall Score", f"{score}%")
    with col2:
        st.metric("Skills Found", len(found_overall))
    with col3:
        st.metric("Missing Skills", len(missing_overall))

    if score >= 80:
        st.success("Excellent Resume — strong alignment with target skills")
    elif score >= 60:
        st.warning("Good Resume — add a few more relevant skills to improve your ATS match")
    else:
        st.error("Resume Needs Improvement — add more key skills and stronger keyword alignment")

    overview, keywords = st.columns(2)
    with overview:
        st.subheader("Resume Summary")
        st.write(f"- Word count: {summary['word_count']}")
        st.write(f"- Email present: {'Yes' if summary['email_present'] else 'No'}")
        st.write(f"- Phone present: {'Yes' if summary['phone_present'] else 'No'}")

    with keywords:
        st.subheader("Top Keywords")
        if top_keywords:
            keyword_text = ", ".join(f"{word} ({count})" for word, count in top_keywords)
            st.info(keyword_text)
        else:
            st.info("No keywords detected")

    st.subheader("Skill Coverage by Category")
    for category, skills in skills_by_category.items():
        category_found = found_by_category[category]
        category_missing = missing_by_category[category]
        with st.expander(f"{category} ({len(category_found)}/{len(skills)})"):
            st.write("Found:")
            st.success(", ".join(category_found) if category_found else "None")
            st.write("Missing:")
            st.error(", ".join(category_missing) if category_missing else "None")

    st.subheader("Recommended Skills to Add")
    recommended = missing_overall[:10]
    if recommended:
        st.warning(", ".join(recommended))
    else:
        st.success("No major skill gaps found")

    st.subheader("Raw Resume Preview")
    st.text_area("Extracted Resume Text", resume_text[:5000], height=250)
